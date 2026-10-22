"""Convert PowerPoint (.pptx) files to Markdown.

For each .pptx in input/, walks every slide with python-pptx, adds a `## Slide N`
heading, promotes large-font text to subheadings, and writes the text to output/.
"""

import io
import re
import shutil
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.shapes.group import GroupShape

import compress
from batch_convert import INPUT_DIR, OUTPUT_DIR, convert_files

# Where python-pptx keeps embedded pictures inside the .pptx zip.
_MEDIA_PREFIX = "ppt/media/"

# These start emphasis, code, links or raw HTML wherever they appear in a line
_INLINE_SPECIALS = re.compile(r'([\\`*\[\]<])')
# An underscore between two letters or digits is literal (snake_case, pptx_md.py). At a
# word edge it can start emphasis. [^\W_] is "a letter or digit".
_EDGE_UNDERSCORE = re.compile(r'(?<![^\W_])_|_(?![^\W_])')
# An ampersand only matters when it looks like an entity (&amp;, &#35;), otherwise R&D is literal
_ENTITY = re.compile(r'&(?=#?[A-Za-z0-9]+;)')


def _escape_inline(text):
    """Backslash-escape the characters that would otherwise be read as Markdown.

    Pass a whole paragraph, cell or heading rather than a single line: emphasis and
    strikethrough can run across a soft line break within a paragraph.
    """
    text = _ENTITY.sub(r'\\&', _INLINE_SPECIALS.sub(r'\\\1', text))
    # Strikethrough needs a pair of tildes, so a lone one (~5 minutes) stays as typed
    if text.count('~') > 1:
        text = text.replace('~', r'\~')
    return _EDGE_UNDERSCORE.sub(r'\\_', text)


def _escape_line_start(line):
    """Escape a marker that would turn the start of a line into a heading, quote or list."""
    line = re.sub(r'^(#{1,6}(?=\s|$)|>|[-+](?=[\s+-]|$))', r'\\\1', line)
    return re.sub(r'^(\d{1,9})([.)])(?=\s|$)', r'\1\\\2', line)


def _cell_text(cell):
    """Cell text on one line, escaped so Markdown characters and pipes stay literal."""
    if cell.is_spanned:
        return ""
    return _escape_inline(" ".join(cell.text.split())).replace("|", "\\|")


def _table_to_markdown(table):
    """Render a table as Markdown, treating the first row as the header.

    Markdown has no column span, so a cell covered by a merge comes out empty and
    the text stays in the cell the merge started from.
    """
    header, *body = [[_cell_text(cell) for cell in row.cells] for row in table.rows]
    lines = ['| ' + ' | '.join(header) + ' |',
             '| ' + ' | '.join(['---'] * len(header)) + ' |']
    lines += ['| ' + ' | '.join(row) + ' |' for row in body]
    return '\n'.join(lines)


def _paragraphs_to_markdown(shape):
    """A shape's paragraphs as Markdown paragraphs, one blank line between them.

    python-pptx returns a soft line break (Shift+Enter) as a raw vertical tab, so it
    becomes a Markdown hard break (a backslash at the end of the line) instead. Inline
    characters are escaped across the whole paragraph, but a line marker is escaped per
    line, since a line after a break can start with one too.
    """
    paragraphs = []
    for paragraph in shape.text.split("\n"):
        lines = [_escape_line_start(line.strip())
                 for line in _escape_inline(paragraph).split("\v") if line.strip()]
        if lines:
            paragraphs.append("\\\n".join(lines))
    return "\n\n".join(paragraphs)


def _append_shape_text(shape, markdown_content):
    """Append one shape's text to markdown_content, recursing into groups.

    Every block ends with a blank line, so neighbouring shapes stay separate
    paragraphs and a table is never glued to the text around it.
    """
    if isinstance(shape, GroupShape):
        for child in shape.shapes:
            _append_shape_text(child, markdown_content)
        return

    if shape.has_table:
        markdown_content.append(f"{_table_to_markdown(shape.table)}\n\n")
        return

    if hasattr(shape, "text") and shape.text.strip():
        # Check if this is likely a title
        if hasattr(shape, "text_frame") and shape.text_frame.paragraphs:
            first_para = shape.text_frame.paragraphs[0]
            if first_para.runs and first_para.runs[0].font.size:
                if first_para.runs[0].font.size > 200000:
                    # A heading is one line, so paragraphs and soft breaks become spaces
                    heading = _escape_inline(' '.join(shape.text.split()))
                    # A trailing run of #s would be read as the heading's closing sequence
                    heading = re.sub(r'(?<![^\s])(#+)$', r'\\\1', heading)
                    markdown_content.append(f"### {heading}\n\n")
                    return

        markdown_content.append(f"{_paragraphs_to_markdown(shape)}\n\n")


def extract_text_from_pptx(pptx_file):
    """Extract text from PowerPoint file and format as markdown"""
    prs = Presentation(pptx_file)
    markdown_content = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        markdown_content.append(f"## Slide {slide_num}\n")

        for shape in slide.shapes:
            _append_shape_text(shape, markdown_content)

        markdown_content.append("---\n\n")

    return ''.join(markdown_content)


def _extract_with_compression(pptx_file: str) -> str:
    """Extract text from `pptx_file`, compressing a throwaway copy first when it has
    images worth downscaling, so a large deck is cheaper to load. The real file in
    input/ is never modified: compression runs on a temp copy, which is discarded
    after extraction either way.
    """
    path = Path(pptx_file)
    if not compress.has_compressible_media(path, _MEDIA_PREFIX):
        return extract_text_from_pptx(pptx_file)

    with tempfile.TemporaryDirectory() as tmp:
        staged = Path(tmp) / path.name
        shutil.copyfile(path, staged)
        compress.compress_office_zip(
            staged, _MEDIA_PREFIX,
            sanity_check=lambda payload: Presentation(io.BytesIO(payload)),
        )
        return extract_text_from_pptx(str(staged))


def _write_markdown(source: Path, dest: Path) -> None:
    # Extract text and format as markdown
    dest.write_text(_extract_with_compression(str(source)), encoding='utf-8')


def convert_pptx_to_markdown(
    input_dir: str | Path = INPUT_DIR, output_dir: str | Path = OUTPUT_DIR
) -> str:
    """Convert all PPTX files in the input folder to Markdown files in the output folder.

    Returns:
        A summary of what was converted, suitable for showing to a caller.
    """
    return convert_files(
        input_dir, output_dir, ('.pptx',), ('.md',), _write_markdown, target_label="markdown"
    )


if __name__ == "__main__":
    result = convert_pptx_to_markdown()
    if not result.startswith("Converted"):
        print(result)
