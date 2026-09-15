"""Convert Word (.docx) documents to PDF.

For each .docx in input/, reads the document with python-docx and rebuilds it as
a PDF with reportlab, carrying over paragraphs, tables, and images. Writes to output/.
"""

import os
import sys
import argparse
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from docx import Document
from docx.oxml.ns import qn
import io

# Default folders live next to this script, so `python backend/docx_pdf.py` finds
# backend/input regardless of the shell's working directory.
SCRIPT_DIR = Path(__file__).resolve().parent


# Create input and output directories if they don't exist
def setup_directories(input_dir="input", output_dir="output"):
    input_path = Path(input_dir)
    output_path = Path(output_dir)
    
    input_path.mkdir(exist_ok=True)
    output_path.mkdir(exist_ok=True)
    
    return input_path, output_path


# convert a docx table to a pdf table and add it to the story
def add_table_to_story(story, table, available_width):
    styles = getSampleStyleSheet()
    header_cell_style = ParagraphStyle(
        'TableHeaderCell', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=10, textColor=colors.whitesmoke,
    )
    body_cell_style = ParagraphStyle(
        'TableBodyCell', parent=styles['Normal'],
        fontName='Helvetica', fontSize=9,
    )

    # Get table data. Cells are wrapped in Paragraphs (rather than left as plain
    # strings) so ReportLab wraps long text to the column width instead of
    # overflowing past the page edge.
    data = []
    for row_index, row in enumerate(table.rows):
        cell_style = header_cell_style if row_index == 0 else body_cell_style
        row_data = []
        for cell in row.cells:
            # Get text from cell, replacing newlines with spaces
            cell_text = cell.text.replace('\n', ' ')
            row_data.append(Paragraph(cell_text, cell_style))
        data.append(row_data)

    if not data:
        return

    # Split the available page width evenly across columns so the table never
    # exceeds the page, whatever column count the original docx table has.
    num_cols = len(data[0])
    col_widths = [available_width / num_cols] * num_cols

    # Create PDF table
    pdf_table = Table(data, colWidths=col_widths)

    # Style the table
    pdf_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
    ]))

    story.append(pdf_table)
    story.append(Spacer(1, 12))


# Convert a .docx file to PDF
# returns True if successful, False otherwise
def convert_docx_to_pdf(docx_path: str, output_path: str | None = None, input_dir: str | None = None, output_dir: str | None = None) -> bool:
    # Set default directories if not provided. Resolved relative to this script, not
    # the caller's working directory, so they are always backend/input and backend/output.
    if input_dir is None:
        input_dir = SCRIPT_DIR / "input"
    else:
        input_dir = Path(input_dir)

    if output_dir is None:
        output_dir = SCRIPT_DIR / "output"
    else:
        output_dir = Path(output_dir)

    # if the provided path already points to a PDF let user know and return False
    try:
        provided_suffix = Path(docx_path).suffix.lower()
        if provided_suffix == ".pdf":
            print("That file is already in pdf format")
            return False
    except Exception:
        pass

    # Build full input path
    full_input_path = Path(docx_path) if os.path.isabs(str(docx_path)) else input_dir / docx_path
    if not full_input_path.exists():
        print(f"Error: Word file '{full_input_path}' not found")
        return False

    # Compute output name
    if output_path is None:
        pdf_name = f"{full_input_path.stem}.pdf"
    else:
        pdf_name = Path(output_path).name
    full_output_path = output_dir / pdf_name

    try:
        # Load the Word document
        docx = Document(str(full_input_path))

        # Create PDF document
        doc = SimpleDocTemplate(str(full_output_path), pagesize=letter)
        styles = getSampleStyleSheet()

        # Create custom styles
        title_style = ParagraphStyle(
            'DocxTitle',
            parent=styles['Title'],
            fontSize=16,
            spaceAfter=20,
            alignment=1  # Center alignment
        )

        normal_style = ParagraphStyle(
            'DocxStyle',
            parent=styles['Normal'],
            fontSize=11,
            leading=14,
            leftIndent=0,
            rightIndent=0,
            spaceAfter=6
        )

        # Build PDF content
        story = []

        # Add title
        filename = full_input_path.stem
        title = Paragraph(f"Document: {filename}", title_style)
        story.append(title)
        story.append(Spacer(1, 20))

        # Process document elements in order (paras and tables interleaved)
        # Create maps for quick lookup
        para_map = {id(para._p): para for para in docx.paragraphs}
        table_map = {id(table._tbl): table for table in docx.tables}
        
        # Iterate through body elements in document order
        for element in docx.element.body:
            elem_id = id(element)
            
            # Check if it's a paragraph
            if elem_id in para_map:
                para = para_map[elem_id]
                
                # Process runs in order to maintain text/image order
                text_runs = []
                has_formatting = False
                
                for run in para.runs:
                    # Check if this run contains an image
                    has_image = False
                    try:
                        drawings = run._element.xpath('.//w:drawing')
                        for drawing in drawings:
                            blips = drawing.xpath('.//a:blip')
                            for blip in blips:
                                rel_id = blip.get(qn('r:embed'))
                                if rel_id and rel_id in docx.part.rels:
                                    has_image = True
                                    rel = docx.part.rels[rel_id]
                                    img_data = rel.target_part.blob
                                    
                                    # Get image dimensions if available
                                    try:
                                        from PIL import Image as PILImage
                                        pil_img = PILImage.open(io.BytesIO(img_data))
                                        img_width, img_height = pil_img.size
                                        # Scale to fit page width (max 5 inches)
                                        max_width = 5 * inch
                                        if img_width > max_width:
                                            ratio = max_width / img_width
                                            img_width = max_width
                                            img_height = img_height * ratio
                                        else:
                                            img_width = img_width * (72 / 96)  # Convert pixels to points
                                            img_height = img_height * (72 / 96)
                                    except Exception:
                                        # Default size if we can't read dimensions
                                        img_width = 5 * inch
                                        img_height = 3 * inch
                                    
                                    img = Image(io.BytesIO(img_data), width=img_width, height=img_height)
                                    story.append(img)
                                    story.append(Spacer(1, 12))
                    except Exception:
                        pass
                    
                    # If run has text and no image, collect it
                    if run.text and not has_image:
                        text_runs.append(run)
                        if run.bold or run.italic or run.underline:
                            has_formatting = True
                
                # Add text from all runs together
                if text_runs:
                    if has_formatting:
                        # Build HTML from all text runs
                        html_parts = []
                        for run in text_runs:
                            text = run.text
                            if not text:
                                continue
                            text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                            if run.bold:
                                text = f'<b>{text}</b>'
                            if run.italic:
                                text = f'<i>{text}</i>'
                            if run.underline:
                                text = f'<u>{text}</u>'
                            html_parts.append(text)
                        if html_parts:
                            story.append(Paragraph(''.join(html_parts), normal_style))
                    else:
                        # Plain text - combine all runs
                        combined_text = ''.join(run.text for run in text_runs)
                        escaped_text = combined_text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        story.append(Paragraph(escaped_text, normal_style))
                elif not any(run._element.xpath('.//w:drawing') for run in para.runs):
                    # Empty paragraph with no images - add spacing
                    story.append(Spacer(1, 6))
            
            # Check if it's a table
            elif elem_id in table_map:
                story.append(Spacer(1, 12))
                add_table_to_story(story, table_map[elem_id], doc.width)

        # Build PDF
        print(f"Converting '{full_input_path}' to '{full_output_path}'...")
        doc.build(story)
        
        if full_output_path.exists():
            print(f"Successfully converted to '{full_output_path}'")
            return True
        else:
            print("PDF creation failed")
            return False

    except Exception as e:
        print(f"Error creating PDF: {e}")
        import traceback
        traceback.print_exc()
        return False


def main():
    parser = argparse.ArgumentParser(description="Convert DOCX files to PDF")
    parser.add_argument("docx_file", nargs="?", help="DOCX file to convert (optional)")
    parser.add_argument("output_file", nargs="?", help="Output PDF filename (optional)")
    parser.add_argument("--input-dir", default=SCRIPT_DIR / "input", help="Input directory (default: backend/input)")
    parser.add_argument("--output-dir", default=SCRIPT_DIR / "output", help="Output directory (default: backend/output)")
    
    args = parser.parse_args()
    
    # Set up directories
    input_dir, output_dir = setup_directories(args.input_dir, args.output_dir)
    
    if args.docx_file:
        # Convert specific file
        success = convert_docx_to_pdf(args.docx_file, args.output_file, input_dir, output_dir)
        if not success:
            sys.exit(1)
    else:
        # No args: convert all .docx files in input/
        docx_files = sorted(p for p in input_dir.glob("*.docx"))
        if not docx_files:
            print(f"No .docx files found in {input_dir} folder")
            print("Usage: python docx_pdf.py <file.docx> [output.pdf] [--input-dir DIR] [--output-dir DIR]")
            print("Example: python docx_pdf.py notes.docx")
            print("Example: python docx_pdf.py notes.docx my_notes.pdf")
            print("Example: python docx_pdf.py --input-dir myinput --output-dir myoutput")
            return
        any_failed = False
        for docx in docx_files:
            ok = convert_docx_to_pdf(docx.name, None, input_dir, output_dir)
            if not ok:
                any_failed = True
        if any_failed:
            sys.exit(1)


if __name__ == "__main__":
    main()
