"""Tests for pptx_md.py.

The decks are built with python-pptx in each test, so there are no binary fixtures.
"""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

import pptx_md

TITLE_ONLY_LAYOUT = 5


def new_slide():
    prs = Presentation()
    return prs, prs.slides.add_slide(prs.slide_layouts[TITLE_ONLY_LAYOUT])


def extract(prs: Presentation, tmp_path: Path) -> str:
    path = tmp_path / "deck.pptx"
    prs.save(path)
    return pptx_md.extract_text_from_pptx(str(path))


def add_textbox(slide, text: str):
    frame = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(3), Inches(1)).text_frame
    frame.text = text
    return frame


def test_each_shape_is_its_own_paragraph(tmp_path: Path) -> None:
    prs, slide = new_slide()
    slide.shapes.title.text = "TITLETEXT"
    add_textbox(slide, "PLAINBOXTEXT")

    assert extract(prs, tmp_path) == "## Slide 1\nTITLETEXT\n\nPLAINBOXTEXT\n\n---\n\n"


def test_paragraphs_in_one_shape_are_separated_and_empty_ones_dropped(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_textbox(slide, "ONE\n\nTWO")

    assert extract(prs, tmp_path) == "## Slide 1\nONE\n\nTWO\n\n---\n\n"


def test_soft_line_break_becomes_a_markdown_hard_break(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_textbox(slide, "before\vafter")

    result = extract(prs, tmp_path)

    assert "before\\\nafter\n\n" in result
    assert "\x0b" not in result


def test_large_font_text_becomes_a_one_line_heading(tmp_path: Path) -> None:
    prs, slide = new_slide()
    frame = add_textbox(slide, "Line one\vLine two\nSecond paragraph")
    frame.paragraphs[0].runs[0].font.size = Pt(24)

    assert "### Line one Line two Second paragraph\n\n" in extract(prs, tmp_path)


@pytest.mark.parametrize("source, expected", [
    ("*star*", "\\*star\\*"),
    ("`code`", "\\`code\\`"),
    ("[a](b)", "\\[a\\](b)"),
    ("<b>", "\\<b>"),
    ("C:\\Users", "C:\\\\Users"),
    ("_lead and trail_", "\\_lead and trail\\_"),
    ("pptx_md.py and snake_case", "pptx_md.py and snake_case"),
    ("# not a heading", "\\# not a heading"),
    ("#hashtag", "#hashtag"),
    ("> not a quote", "\\> not a quote"),
    ("- not a bullet", "\\- not a bullet"),
    ("+ not a bullet", "\\+ not a bullet"),
    ("---", "\\---"),
    ("1. not a list", "1\\. not a list"),
    ("2) not a list", "2\\) not a list"),
    ("3.5 stars", "3.5 stars"),
    ("~~struck~~", "\\~\\~struck\\~\\~"),
    ("about ~5 minutes", "about ~5 minutes"),
    ("&amp; &#35; &copy;", "\\&amp; \\&#35; \\&copy;"),
    ("R&D and AT&T", "R&D and AT&T"),
])
def test_markdown_characters_in_text_stay_literal(tmp_path: Path, source: str, expected: str) -> None:
    prs, slide = new_slide()
    add_textbox(slide, source)

    assert extract(prs, tmp_path) == f"## Slide 1\n{expected}\n\n---\n\n"


def test_a_marker_after_a_soft_line_break_is_escaped(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_textbox(slide, "before\v# after")

    assert "before\\\n\\# after\n\n" in extract(prs, tmp_path)


def test_a_tilde_pair_split_by_a_soft_line_break_is_escaped(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_textbox(slide, "~one\vtwo~")

    assert "\\~one\\\ntwo\\~\n\n" in extract(prs, tmp_path)


def test_markdown_characters_in_a_heading_stay_literal(tmp_path: Path) -> None:
    prs, slide = new_slide()
    frame = add_textbox(slide, "*Big* title")
    frame.paragraphs[0].runs[0].font.size = Pt(24)

    assert "### \\*Big\\* title\n\n" in extract(prs, tmp_path)


@pytest.mark.parametrize("source, expected", [
    ("Title #", "Title \\#"),
    ("Title ##", "Title \\##"),
    ("#", "\\#"),
    ("C#", "C#"),
])
def test_a_trailing_hash_in_a_heading_is_not_a_closing_sequence(
    tmp_path: Path, source: str, expected: str
) -> None:
    prs, slide = new_slide()
    frame = add_textbox(slide, source)
    frame.paragraphs[0].runs[0].font.size = Pt(24)

    assert f"### {expected}\n\n" in extract(prs, tmp_path)


def add_table(slide, cells: list[list[str]]):
    rows, cols = len(cells), len(cells[0])
    table = slide.shapes.add_table(rows, cols, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for r, row in enumerate(cells):
        for c, text in enumerate(row):
            table.cell(r, c).text = text
    return table


def test_table_becomes_a_markdown_table(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_table(slide, [["A", "B"], ["C", "D"]])

    assert extract(prs, tmp_path) == "## Slide 1\n| A | B |\n| --- | --- |\n| C | D |\n\n---\n\n"


def test_text_after_a_table_is_not_glued_to_it(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_table(slide, [["A", "B"], ["C", "D"]])
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "AFTER"

    assert "| C | D |\n\nAFTER" in extract(prs, tmp_path)


def test_pipes_and_line_breaks_stay_inside_their_cell(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_table(slide, [["a|b", "one\ntwo"], ["x", "y"]])

    assert "| a\\|b | one two |\n| --- | --- |\n| x | y |" in extract(prs, tmp_path)


def test_markdown_characters_in_a_cell_stay_literal(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_table(slide, [["*a* | b", "snake_case"], ["1. x", "y"]])

    assert "| \\*a\\* \\| b | snake_case |\n| --- | --- |\n| 1. x | y |" in extract(prs, tmp_path)


def test_merged_cell_keeps_its_text_and_the_column_count(tmp_path: Path) -> None:
    prs, slide = new_slide()
    table = add_table(slide, [["", ""], ["A", "B"]])
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "MERGED"

    assert "| MERGED |  |\n| --- | --- |\n| A | B |" in extract(prs, tmp_path)


def test_empty_first_row_still_gets_a_header(tmp_path: Path) -> None:
    prs, slide = new_slide()
    add_table(slide, [["", ""], ["A", "B"]])

    assert "|  |  |\n| --- | --- |\n| A | B |" in extract(prs, tmp_path)


def test_text_inside_a_group_is_kept(tmp_path: Path) -> None:
    prs, slide = new_slide()
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(5), Inches(1.5), Inches(3), Inches(1)).text_frame.text = "GROUPTEXT"

    assert "GROUPTEXT" in extract(prs, tmp_path)


def test_text_inside_a_nested_group_is_kept(tmp_path: Path) -> None:
    prs, slide = new_slide()
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(5), Inches(1.5), Inches(3), Inches(1)).text_frame.text = "NESTEDTEXT"

    assert "NESTEDTEXT" in extract(prs, tmp_path)
