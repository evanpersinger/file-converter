"""Fixtures shared by the converter tests.

The converters take their input and output folders as arguments, so a test hands them a
tmp_path pair (`work_dirs`) and never touches the real backend/input and backend/output.
"""

import shutil
from collections.abc import Callable
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from external_tool import run_tool


@pytest.fixture
def work_dirs(tmp_path: Path) -> tuple[Path, Path]:
    """A throwaway (input_dir, output_dir) pair to hand to a converter.

    Only the input directory is created, so a converter making its own output directory
    stays testable.
    """
    input_dir = tmp_path / "input"
    input_dir.mkdir()
    return input_dir, tmp_path / "output"


@pytest.fixture
def write_image() -> Callable[..., Path]:
    """Write a small genuine image, so Pillow has something real to open.

    The format is passed explicitly rather than inferred, so a test can deliberately
    mismatch the contents and the extension.
    """

    def write(path: Path, fmt: str = "JPEG", color: str = "red") -> Path:
        Image.new("RGB", (8, 8), color).save(path, fmt)
        return path

    return write


@pytest.fixture
def write_ppt(tmp_path: Path) -> Callable[..., Path]:
    """Write a genuine old-format .ppt with one slide titled `title`.

    python-pptx only writes .pptx, so LibreOffice saves it down to .ppt. Only for tests
    marked as needing LibreOffice.
    """

    def write(path: Path, title: str) -> Path:
        deck = Presentation()
        deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = title
        source = tmp_path / "ppt_source" / f"{path.stem}.pptx"
        source.parent.mkdir(exist_ok=True)
        deck.save(source)
        run_tool([
            shutil.which("soffice") or "libreoffice", "--headless",
            f"-env:UserInstallation={(tmp_path / 'ppt_profile').as_uri()}",
            "--convert-to", "ppt", "--outdir", str(path.parent), str(source),
        ])
        assert path.exists()
        return path

    return write
