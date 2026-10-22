"""Tests for docx_md.py.

These run the real pandoc, so they need it installed. The image test matters most:
pandoc writes image links exactly as it is given the media path, so a regression there
produces Markdown whose images break the moment the output is zipped or moved.
"""

from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches

import docx_md

WorkDirs = tuple[Path, Path]


def test_reports_when_input_is_empty(work_dirs: WorkDirs) -> None:
    input_dir, output_dir = work_dirs
    assert (
        docx_md.convert_docx_to_markdown(input_dir, output_dir)
        == "No DOCX files found in input folder"
    )


def test_reports_when_pandoc_is_missing(
    work_dirs: WorkDirs, monkeypatch: pytest.MonkeyPatch
) -> None:
    input_dir, output_dir = work_dirs
    monkeypatch.setattr(docx_md, "which", lambda _name: None)

    assert docx_md.convert_docx_to_markdown(input_dir, output_dir).startswith(
        "Error: pandoc not found"
    )


def test_keeps_headings_formatting_and_tables(work_dirs: WorkDirs) -> None:
    input_dir, output_dir = work_dirs
    doc = Document()
    doc.add_heading("Notes", level=1)
    paragraph = doc.add_paragraph("This is ")
    paragraph.add_run("bold").bold = True
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Evan"
    table.cell(1, 1).text = "92"
    doc.save(input_dir / "notes.docx")

    summary = docx_md.convert_docx_to_markdown(input_dir, output_dir)

    assert "notes.md" in summary
    markdown = (output_dir / "notes.md").read_text(encoding="utf-8")
    assert "# Notes" in markdown
    assert "**bold**" in markdown
    assert "| Name | Score |" in markdown


def test_images_are_saved_and_linked_with_relative_paths(
    work_dirs: WorkDirs, write_image: Callable[..., Path], tmp_path: Path
) -> None:
    input_dir, output_dir = work_dirs
    doc = Document()
    doc.add_picture(str(write_image(tmp_path / "red.png", "PNG")), width=Inches(1))
    doc.save(input_dir / "report.docx")

    docx_md.convert_docx_to_markdown(input_dir, output_dir)

    markdown = (output_dir / "report.md").read_text(encoding="utf-8")
    assert "report_images/media/image1.png" in markdown
    assert str(tmp_path) not in markdown
    assert (output_dir / "report_images" / "media" / "image1.png").exists()
