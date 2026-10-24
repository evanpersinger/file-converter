"""Tests for docx_pdf.py, which rebuilds a Word document as a PDF with reportlab.

It has its own loop rather than convert_files: it returns True or False and prints.
"""

from collections.abc import Callable
from pathlib import Path

import fitz
import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches

import docx_pdf

WorkDirs = tuple[Path, Path]
WriteImage = Callable[..., Path]


def convert(work_dirs: WorkDirs, document: Document) -> Path:
    """Save the document as report.docx, convert it, and return the PDF path."""
    input_dir, output_dir = work_dirs
    output_dir.mkdir(exist_ok=True)
    document.save(input_dir / "report.docx")
    assert docx_pdf.convert_docx_to_pdf("report.docx", None, input_dir, output_dir) is True
    return output_dir / "report.pdf"


def pdf_text(path: Path) -> str:
    with fitz.open(path) as pdf:
        return "".join(page.get_text() for page in pdf)


def test_paragraphs_and_the_title_reach_the_pdf(work_dirs: WorkDirs) -> None:
    document = Document()
    document.add_paragraph("first paragraph")
    document.add_paragraph("second paragraph")

    text = pdf_text(convert(work_dirs, document))

    assert "Document: report" in text
    assert "first paragraph" in text and "second paragraph" in text


def test_bold_and_italic_runs_are_styled(work_dirs: WorkDirs) -> None:
    document = Document()
    paragraph = document.add_paragraph("plain ")
    paragraph.add_run("heavy ").bold = True
    paragraph.add_run("slanted").italic = True

    with fitz.open(convert(work_dirs, document)) as pdf:
        fonts = {
            span["text"].strip(): span["font"]
            for block in pdf[0].get_text("dict")["blocks"]
            for line in block.get("lines", [])
            for span in line["spans"]
        }

    assert "Bold" in fonts["heavy"]
    assert "Oblique" in fonts["slanted"]


def test_text_inside_a_hyperlink_is_kept(work_dirs: WorkDirs) -> None:
    """python-docx's paragraph.runs skips runs nested in a hyperlink element."""
    document = Document()
    paragraph = document.add_paragraph("See ")
    hyperlink = OxmlElement("w:hyperlink")
    run = OxmlElement("w:r")
    run_text = OxmlElement("w:t")
    run_text.text = "the docs"
    run.append(run_text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)

    assert "the docs" in pdf_text(convert(work_dirs, document))


def test_a_table_keeps_its_cells(work_dirs: WorkDirs) -> None:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    for (row, column), value in {(0, 0): "Name", (0, 1): "Score", (1, 0): "Evan", (1, 1): "92"}.items():
        table.cell(row, column).text = value

    text = pdf_text(convert(work_dirs, document))

    assert all(cell in text for cell in ("Name", "Score", "Evan", "92"))


def test_an_embedded_image_is_drawn(
    work_dirs: WorkDirs, write_image: WriteImage, tmp_path: Path
) -> None:
    document = Document()
    document.add_picture(str(write_image(tmp_path / "red.png", "PNG")), width=Inches(1))

    with fitz.open(convert(work_dirs, document)) as pdf:
        assert len(pdf[0].get_images()) == 1


def test_markup_characters_are_printed_as_text(work_dirs: WorkDirs) -> None:
    document = Document()
    document.add_paragraph("1 < 2 & 3 > 2")

    assert "1 < 2 & 3 > 2" in pdf_text(convert(work_dirs, document))


def test_a_pdf_path_is_refused(
    work_dirs: WorkDirs, capsys: pytest.CaptureFixture[str]
) -> None:
    input_dir, output_dir = work_dirs

    assert docx_pdf.convert_docx_to_pdf("doc.pdf", None, input_dir, output_dir) is False
    assert "already in pdf format" in capsys.readouterr().out


def test_a_missing_file_gives_false(
    work_dirs: WorkDirs, capsys: pytest.CaptureFixture[str]
) -> None:
    input_dir, output_dir = work_dirs

    assert docx_pdf.convert_docx_to_pdf("ghost.docx", None, input_dir, output_dir) is False
    assert "not found" in capsys.readouterr().out


def test_an_existing_pdf_is_overwritten(work_dirs: WorkDirs) -> None:
    _, output_dir = work_dirs
    output_dir.mkdir()
    (output_dir / "report.pdf").write_bytes(b"stale")
    document = Document()
    document.add_paragraph("fresh text")

    assert "fresh text" in pdf_text(convert(work_dirs, document))
