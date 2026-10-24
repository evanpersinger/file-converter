"""Tests for pptx_pdf.py, which has headless LibreOffice render each PPTX to a PDF.

LibreOffice is faked by patching run_tool, so the command and the error handling are exact;
one case runs the real thing and is skipped without it.
"""

import shutil
import subprocess
from collections.abc import Callable
from pathlib import Path

import fitz
import pytest
from pptx import Presentation

import pptx_pdf

WorkDirs = tuple[Path, Path]

needs_libreoffice = pytest.mark.skipif(
    not any(shutil.which(name) for name in ("libreoffice", "soffice")),
    reason="LibreOffice is not on PATH",
)


def fake_libreoffice(
    commands: list[list[str]], returncode: int = 0, stderr: str = ""
) -> Callable[..., subprocess.CompletedProcess]:
    """A run_tool that records the command and, like LibreOffice, writes the PDF to --outdir."""

    def run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess:
        commands.append(command)
        if returncode == 0:
            outdir = Path(command[command.index("--outdir") + 1])
            (outdir / f"{Path(command[-1]).stem}.pdf").write_bytes(b"%PDF-1.4")
        return subprocess.CompletedProcess(command, returncode, "", stderr)

    return run


def test_the_first_libreoffice_that_runs_is_used(monkeypatch: pytest.MonkeyPatch) -> None:
    def run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess:
        if command[0] == "libreoffice":
            raise FileNotFoundError(command[0])
        return subprocess.CompletedProcess(command, 0, "LibreOffice 24.2", "")

    monkeypatch.setattr(pptx_pdf, "run_tool", run)

    assert pptx_pdf._find_libreoffice() == "soffice"


def test_no_libreoffice_gives_none(monkeypatch: pytest.MonkeyPatch) -> None:
    def run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess:
        raise FileNotFoundError(command[0])

    monkeypatch.setattr(pptx_pdf, "run_tool", run)

    assert pptx_pdf._find_libreoffice() is None


def test_a_missing_libreoffice_is_reported_with_install_help(
    work_dirs: WorkDirs, monkeypatch: pytest.MonkeyPatch
) -> None:
    input_dir, output_dir = work_dirs
    (input_dir / "deck.pptx").write_bytes(b"placeholder")
    monkeypatch.setattr(pptx_pdf, "_find_libreoffice", lambda: None)

    summary = pptx_pdf.convert_pptx_to_pdf(input_dir, output_dir)

    assert summary.startswith("LibreOffice not found in PATH")
    assert "brew install --cask libreoffice" in summary
    assert list(output_dir.iterdir()) == []


def test_a_failed_libreoffice_run_is_reported_with_its_error(
    work_dirs: WorkDirs, monkeypatch: pytest.MonkeyPatch
) -> None:
    input_dir, output_dir = work_dirs
    (input_dir / "deck.pptx").write_bytes(b"placeholder")
    monkeypatch.setattr(pptx_pdf, "_find_libreoffice", lambda: "soffice")
    monkeypatch.setattr(
        pptx_pdf, "run_tool", fake_libreoffice([], returncode=1, stderr="source file could not be loaded")
    )

    summary = pptx_pdf.convert_pptx_to_pdf(input_dir, output_dir)

    assert summary == "No files converted. 1 failed: deck.pptx: source file could not be loaded"


def test_libreoffice_is_told_to_write_into_the_output_folder(
    work_dirs: WorkDirs, monkeypatch: pytest.MonkeyPatch
) -> None:
    input_dir, output_dir = work_dirs
    (input_dir / "deck.pptx").write_bytes(b"placeholder")
    commands: list[list[str]] = []
    monkeypatch.setattr(pptx_pdf, "_find_libreoffice", lambda: "soffice")
    monkeypatch.setattr(pptx_pdf, "run_tool", fake_libreoffice(commands))

    summary = pptx_pdf.convert_pptx_to_pdf(input_dir, output_dir)

    command = commands[0]
    assert command[0] == "soffice"
    assert "--headless" in command
    assert command[command.index("--convert-to") + 1] == "pdf"
    assert command[command.index("--outdir") + 1] == str(output_dir)
    assert command[-1] == str(input_dir / "deck.pptx")
    assert "deck.pdf" in summary


def test_each_call_gets_its_own_libreoffice_profile(
    work_dirs: WorkDirs, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Two runs sharing LibreOffice's default profile collide, and one of them fails."""
    input_dir, output_dir = work_dirs
    (input_dir / "deck.pptx").write_bytes(b"placeholder")
    commands: list[list[str]] = []
    monkeypatch.setattr(pptx_pdf, "_find_libreoffice", lambda: "soffice")
    monkeypatch.setattr(pptx_pdf, "run_tool", fake_libreoffice(commands))

    pptx_pdf.convert_pptx_to_pdf(input_dir, output_dir)
    pptx_pdf.convert_pptx_to_pdf(input_dir, output_dir)

    profiles = [
        argument for command in commands for argument in command
        if argument.startswith("-env:UserInstallation=file://")
    ]
    assert len(profiles) == 2
    assert profiles[0] != profiles[1]


@needs_libreoffice
def test_real_libreoffice_renders_the_slide(work_dirs: WorkDirs) -> None:
    input_dir, output_dir = work_dirs
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Hello deck"
    deck.save(input_dir / "deck.pptx")

    summary = pptx_pdf.convert_pptx_to_pdf(input_dir, output_dir)

    assert "deck.pdf" in summary
    with fitz.open(output_dir / "deck.pdf") as pdf:
        assert pdf.page_count == 1
        assert "Hello deck" in pdf[0].get_text()
